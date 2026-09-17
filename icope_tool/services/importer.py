"""從 Excel 匯入轉介資源。

支援兩種格式（依表頭自動判斷）：
1. 衛生局「社區資源盤點表」：序號／分類名稱／個別名稱／型態／類別／地址／電話／電郵／官網
   「類別」欄的次類別字母 A–F 對應認知、行動、營養、視力、聽力、憂鬱。
   （轉換邏輯移植自參考專案的 scripts/convert_community_resources.py）
2. 本工具的範本（「下載 Excel 範本」產生）：名稱／類型／適用項目／電話／地址／聯絡人／服務時間／網站／備註／常用
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from icope_tool.domains import DOMAIN_BY_EXCEL_CODE, DOMAIN_BY_NAME, DOMAIN_IDS, DOMAINS
from icope_tool.models import Resource

TEMPLATE_HEADERS = ["名稱", "類型", "適用項目", "電話", "地址", "聯絡人", "服務時間", "網站", "備註", "常用"]
_INVENTORY_KEYS = {
    "seq": ("序號",), "category": ("分類名稱", "分類"), "name": ("個別名稱",),
    "class": ("類別",), "address": ("地址",), "phone": ("電話",), "website": ("官網", "網站"),
}
_TEMPLATE_KEYS = {
    "name": ("名稱",), "type": ("類型",), "domains": ("適用項目",), "phone": ("電話",),
    "address": ("地址",), "contact": ("聯絡人",), "hours": ("服務時間", "時間"),
    "website": ("網站",), "note": ("備註",), "pinned": ("常用",),
}
_HEADER_SCAN_ROWS = 30


class ImportError_(Exception):
    """無法解析的 Excel（訊息可直接顯示）。"""


@dataclass
class ImportPreview:
    format: str                       # inventory / template
    sheet: str
    resources: list[Resource] = field(default_factory=list)
    skipped_rows: int = 0
    warnings: list[str] = field(default_factory=list)

    @property
    def format_label(self) -> str:
        return "社區資源盤點表" if self.format == "inventory" else "本工具的匯入範本"


def _clean(value) -> str | None:
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value)).strip()
    if text in ("", "-", "—", "無"):
        return None
    return text


def _header_key(value) -> str:
    return re.sub(r"[\s*＊]", "", str(value or ""))


def _map_columns(row, keys: dict[str, tuple[str, ...]]) -> dict[str, int]:
    cells = [_header_key(v) for v in row]
    mapping = {}
    for key, names in keys.items():
        for index, cell in enumerate(cells):
            if cell in names:
                mapping[key] = index
                break
    return mapping


def split_phones(text: str | None) -> list[str]:
    if not text:
        return []
    parts = re.split(r"[、,，;；/]|\s{2,}|\s(?=0\d)", text)
    return [p.strip() for p in parts if p and p.strip()]


def parse_domain_text(text: str | None) -> tuple[list[str], list[str]]:
    """「認知、行動」→ ([cognitive, mobility], [])；回傳 (domain ids, 看不懂的片段)。"""
    if not text:
        return [], []
    if text.strip() in ("全部", "全", "所有"):
        return list(DOMAIN_IDS), []
    ids, unknown = [], []
    for part in re.split(r"[、,，;；/\s]+", text):
        part = part.strip()
        if not part:
            continue
        domain = DOMAIN_BY_NAME.get(part) or DOMAIN_BY_NAME.get(part.replace("功能", "").replace("能力", ""))
        if domain:
            ids.append(domain.id)
        else:
            unknown.append(part)
    return ids, unknown


def list_sheets(path: Path) -> list[str]:
    import openpyxl
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ImportError_("無法開啟這個 Excel 檔，請確認是 .xlsx 格式且沒有被其他程式開著") from exc
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def parse_resources_xlsx(path: Path, sheet: str | None = None) -> ImportPreview:
    import openpyxl
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ImportError_("無法開啟這個 Excel 檔，請確認是 .xlsx 格式且沒有被其他程式開著") from exc
    try:
        names = list(workbook.sheetnames)
        if sheet is None:
            sheet = _guess_sheet(workbook, names)
        if sheet not in names:
            raise ImportError_(f"找不到工作表「{sheet}」")
        rows = list(workbook[sheet].iter_rows(values_only=True))
    finally:
        workbook.close()

    for index, row in enumerate(rows[:_HEADER_SCAN_ROWS]):
        keys = {_header_key(v) for v in row}
        if "個別名稱" in keys:
            return _parse_inventory(sheet, rows, index)
        if "名稱" in keys and "適用項目" in keys:
            return _parse_template(sheet, rows, index)
    raise ImportError_(
        f"工作表「{sheet}」找不到可辨識的表頭。請使用「社區資源盤點表」格式，"
        "或按「下載 Excel 範本」依範本填寫後再匯入。"
    )


def _guess_sheet(workbook, names: list[str]) -> str:
    """盤點表慣例：最新資料在最後一個工作表；範本則是第一個。找第一個有可辨識表頭的（由後往前）。"""
    for name in reversed(names):
        for row in workbook[name].iter_rows(values_only=True, max_row=_HEADER_SCAN_ROWS):
            keys = {_header_key(v) for v in row}
            if "個別名稱" in keys or ("名稱" in keys and "適用項目" in keys):
                return name
    return names[-1]


def _cell(row, mapping, key):
    index = mapping.get(key)
    if index is None or index >= len(row):
        return None
    return _clean(row[index])


def _parse_inventory(sheet: str, rows, header_index: int) -> ImportPreview:
    mapping = _map_columns(rows[header_index], _INVENTORY_KEYS)
    preview = ImportPreview("inventory", sheet)
    started = False
    last_category = None
    for row in rows[header_index + 1:]:
        seq = _cell(row, mapping, "seq")
        name = _cell(row, mapping, "name")
        if not started:
            # 表頭下常有外縣市範例列，沒有數字序號；遇到第一筆數字序號才開始
            if seq and seq.isdigit():
                started = True
            else:
                continue
        if not name:
            if any(_clean(v) for v in row):
                preview.skipped_rows += 1
            continue
        category = _cell(row, mapping, "category") or last_category or "其他"
        last_category = category
        class_code = (_cell(row, mapping, "class") or "").upper()
        domains = [DOMAIN_BY_EXCEL_CODE[c].id for c in class_code if c in DOMAIN_BY_EXCEL_CODE]
        preview.resources.append(Resource(
            name=name, type=category, domains=domains,
            phones=split_phones(_cell(row, mapping, "phone")),
            address=_cell(row, mapping, "address") or "",
            website=_cell(row, mapping, "website") or "",
        ))
    if not preview.resources:
        preview.warnings.append("沒有讀到任何資源。盤點表需要「序號」欄為數字的資料列。")
    untagged = sum(1 for r in preview.resources if not r.domains)
    if untagged:
        preview.warnings.append(
            f"有 {untagged} 筆在「類別」欄沒有 A–F 次類別，不會出現在任何評估項目的建議清單，"
            "但仍可在選轉介時用搜尋加入；也可以匯入後到資源庫補標適用項目。"
        )
    return preview


def _parse_template(sheet: str, rows, header_index: int) -> ImportPreview:
    mapping = _map_columns(rows[header_index], _TEMPLATE_KEYS)
    preview = ImportPreview("template", sheet)
    unknown_parts: set[str] = set()
    for number, row in enumerate(rows[header_index + 1:], start=header_index + 2):
        name = _cell(row, mapping, "name")
        if not name:
            if any(_clean(v) for v in row):
                preview.skipped_rows += 1
                preview.warnings.append(f"第 {number} 列沒有填名稱，已略過")
            continue
        if name.startswith("範例："):
            continue
        domains, unknown = parse_domain_text(_cell(row, mapping, "domains"))
        unknown_parts.update(unknown)
        pinned_text = _cell(row, mapping, "pinned") or ""
        preview.resources.append(Resource(
            name=name, type=_cell(row, mapping, "type") or "", domains=domains,
            phones=split_phones(_cell(row, mapping, "phone")),
            address=_cell(row, mapping, "address") or "", contact=_cell(row, mapping, "contact") or "",
            hours=_cell(row, mapping, "hours") or "", website=_cell(row, mapping, "website") or "",
            note=_cell(row, mapping, "note") or "",
            pinned=pinned_text in ("是", "Y", "y", "V", "v", "✓", "1", "常用"),
        ))
    if unknown_parts:
        valid = "、".join(d.name for d in DOMAINS)
        preview.warnings.append(f"「適用項目」看不懂：{'、'.join(sorted(unknown_parts))}（可填：{valid}，或「全部」）")
    return preview


def write_template_xlsx(path: Path) -> None:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "轉介資源"
    sheet.append(TEMPLATE_HEADERS)
    sheet.append(["範例：嘉義基督教醫院 記憶門診", "醫療院所", "認知", "05-2765041",
                  "嘉義市東區忠孝路539號", "", "週一、三、五上午", "", "需先掛號", "是"])
    header_fill = PatternFill("solid", fgColor="0F766E")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    for cell in sheet[2]:
        cell.font = Font(color="64748B", italic=True)
    widths = [34, 16, 18, 22, 36, 12, 18, 28, 24, 8]
    for column, width in zip("ABCDEFGHIJ", widths):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"

    guide = workbook.create_sheet("填寫說明")
    lines = [
        ("填寫說明", True),
        ("1. 每一列一個轉介資源；「名稱」必填，其他欄位可留空。", False),
        ("2. 「範例：」開頭的列不會匯入，可以直接刪除或保留。", False),
        ("3. 「適用項目」填評估項目名稱，多個用頓號分隔；全部都適用可填「全部」。", False),
        ("   可填：" + "、".join(d.name for d in DOMAINS), False),
        ("4. 「電話」有多支時用頓號分隔，例如 05-2765041、0912345678。", False),
        ("5. 「常用」填「是」會在選轉介時排在最上面。", False),
        ("6. 名稱＋地址和現有資源相同的列，匯入時會自動略過，不會重複。", False),
    ]
    for text, bold in lines:
        guide.append([text])
        guide.cell(row=guide.max_row, column=1).font = Font(bold=bold, size=12 if bold else 11)
    guide.column_dimensions["A"].width = 90
    workbook.save(path)
