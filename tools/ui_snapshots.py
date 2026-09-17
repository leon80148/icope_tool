"""產生 UI 審查用截圖（示範資料＋假的國健署回應，不連網、不碰讀卡機）。

用法：python tools/ui_snapshots.py <輸出資料夾> [--size 1366x768] [--only 檔名片段]

檔名開頭：q 篩檢查詢、r 轉介衛教單、p 印出的 PDF、s 設定、d 對話框、f 鍵盤焦點（*_crop 為放大局部）。
"""
from __future__ import annotations

import argparse
import io
import os
import shutil
import sys
import tempfile
import threading
import time
import warnings
from datetime import datetime, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

DEMO_IDS = {"王小明": "A123456789", "林阿土": "B123456780", "陳秀英": "F223456786", "張美玉": "C223456783"}


def make_material_pdf(path: Path, title: str, pages: int) -> Path:
    from fpdf import FPDF
    from icope_tool.services.pdf.referral_sheet import FONT_DIR
    pdf = FPDF()
    pdf.add_font("tc", "", str(FONT_DIR / "NotoSansTC-Regular.ttf"))
    for page in range(pages):
        pdf.add_page()
        pdf.set_font("tc", size=26)
        pdf.cell(text=f"{title}")
        pdf.ln(16)
        pdf.set_font("tc", size=14)
        pdf.multi_cell(0, 9, "這是示範用的衛教單張內容。" * 8)
        pdf.ln(4)
        pdf.cell(text=f"第 {page + 1} 頁")
    pdf.output(str(path))
    return path


def captcha_gif() -> bytes:
    from PIL import Image, ImageDraw, ImageFont
    image = Image.new("RGB", (100, 34), (240, 244, 248))
    draw = ImageDraw.Draw(image)
    for x in range(0, 100, 9):
        draw.line([(x, 0), (x + 20, 34)], fill=(190, 200, 220))
    try:
        font = ImageFont.truetype("arial.ttf", 22)
    except OSError:
        font = ImageFont.load_default()
    for index, digit in enumerate("48157"):
        draw.text((8 + index * 18, 5), digit, fill=(40 + index * 30, 60, 160), font=font)
    buffer = io.BytesIO()
    image.save(buffer, format="GIF")
    return buffer.getvalue()


class FakeHpdcs:
    def __init__(self, configured: bool = True):
        self.configured = configured
        self.logged_in = configured
        self.disabled = False

    def status(self):
        return {"configured": self.configured, "account_masked": "dem***" if self.configured else "",
                "logged_in": self.logged_in, "auto_login_disabled": self.disabled,
                "last_login_at": datetime.now().isoformat(timespec="seconds") if self.logged_in else None,
                "last_error": None, "ocr_available": True,
                "active_plans": [{"plan": "EFA_115", "label": "115 年度正式計畫"},
                                 {"plan": "EFA_Pilot_115", "label": "115 年度試辦計畫"}]}

    def query_icope(self, *args, **kwargs):  # noqa: D401
        raise RuntimeError("snapshots drive the UI directly")

    def login(self, *args, **kwargs):
        return None

    def refresh_captcha(self, token):
        from icope_tool.services.hpdcs.client import CaptchaManualRequired
        return CaptchaManualRequired(captcha_gif(), "image/gif", token + "x")

    def invalidate(self):
        pass

    def clear_cache(self):
        pass

    def warm_up_ocr(self):
        return True


def make_demo_store(work: Path):
    """示範資料（虛構的診所與長者、公開的嘉義市資源、程式產生的假單張）。截圖與說明圖片共用，不含任何真實個資。"""
    from icope_tool.models import Resource
    from icope_tool.services.pack import apply_pack, read_pack
    from icope_tool.store import DataStore

    store = DataStore(work / "data")
    store.ensure_ready(create=True)
    store.update_settings(lambda s: (setattr(s.clinic, "name", "示範診所"), setattr(s.clinic, "phone", "05-0000000"),
                                     setattr(s.clinic, "address", "嘉義市東區示範路 100 號"),
                                     setattr(s.clinic, "footer_note", "本單僅供轉介參考，服務內容以各單位公告為準")))
    apply_pack(store, read_pack(ROOT / "examples" / "chiayi-city-resources.json"), "merge")
    store.update_settings(lambda s: s.sheet.domain_notes.update({
        "cognitive": "多動腦、多互動：每天和家人聊聊今天做了什麼，一起看老照片、玩簡單的牌卡或算數遊戲。\n忘東忘西變多、迷路或個性改變時，請提早回診。",
        "mobility": "每天做腿部肌力與平衡運動（扶著椅背練習坐下、站起 10 次）；走道保持明亮、移開雜物，浴室加裝扶手，預防跌倒。"}))
    store.add_resource(Resource(name="嘉義基督教醫院 記憶門診", type="醫療院所", pinned=True, include_by_default=True,
                                domains=["cognitive"], phones=["05-2765041"], hours="週一、三、五上午",
                                address="嘉義市東區忠孝路539號", note="需先掛號"))
    store.add_resource(Resource(name="嘉義市政府衛生局附設長期照顧管理中心（嘉義市長照管理中心東區辦公室）",
                                type="長照管理中心", domains=["mobility", "nutrition", "social"], include_by_default=True,
                                phones=["05-2336889", "1966"], hours="週一至週五 8:00–17:00",
                                address="嘉義市東區德明路 1 號"))
    pdf_dir = work / "pdfs"
    pdf_dir.mkdir()
    for name, pages, domains, desc in (("銀髮族營養指南", 2, ["nutrition"], "一日飲食建議與蛋白質補充"),
                                       ("預防跌倒與居家安全", 1, ["mobility"], "居家環境檢查表"),
                                       ("認知促進活動", 2, ["cognitive", "social"], ""),
                                       ("用藥安全五大核心", 1, ["medication"], "")):
        store.add_material_file(make_material_pdf(pdf_dir / f"{name}.pdf", name, pages), name, desc, domains,
                                include_by_default=name in ("預防跌倒與居家安全", "認知促進活動"))
    (store.auth_dir / "hpdcs_credentials.json").write_text('{"account":"demo001","password":"x"}', encoding="utf-8")
    return store


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("out", type=Path)
    parser.add_argument("--size", default="1366x768")
    parser.add_argument("--only", default="", help="只輸出檔名含此字串的截圖")
    args = parser.parse_args()
    width, height = (int(v) for v in args.size.lower().split("x"))
    out: Path = args.out
    out.mkdir(parents=True, exist_ok=True)

    from PySide6.QtCore import QPoint, QRect, QSize, Qt
    from PySide6.QtWidgets import QApplication, QRadioButton, QWidget

    app = QApplication(sys.argv[:1])
    from icope_tool.audit import AuditLog
    from icope_tool.models import Resource
    from icope_tool.services.hpdcs.client import CaptchaManualRequired, CredentialError, IcopeResult, PlanResult
    from icope_tool.services.pack import apply_pack, read_pack
    from icope_tool.store import DataStore, LocalConfigStore
    from icope_tool.ui.context import AppContext, HistoryEntry
    from icope_tool.ui.main_window import MainWindow
    from icope_tool.ui.query_page import QueryRequest
    from icope_tool.ui.theme import apply_theme
    from icope_tool.ui.workers import keep_garbage_collection_on_main_thread
    apply_theme(app)
    gc_timer = keep_garbage_collection_on_main_thread(app)   # noqa: F841 - 與正式程式相同，避免背景回收卡死

    work = Path(tempfile.mkdtemp(prefix="icope-snap-"))

    def pump(seconds: float = 0.15):
        end = time.monotonic() + seconds
        while time.monotonic() < end:
            app.processEvents()
            time.sleep(0.01)

    def wanted(name: str) -> bool:
        return not args.only or args.only in name

    def open_window(ctx) -> MainWindow:
        window = MainWindow(ctx)
        window.setAttribute(Qt.WidgetAttribute.WA_DontShowOnScreen, True)
        window.resize(width, height)
        window.show()
        return window

    def shot(name: str, target: QWidget):
        if not wanted(name):
            return
        pump(0.6)
        target.grab().save(str(out / f"{name}.png"))
        print("saved", name)

    def dialog_shot(name: str, dialog, ready=None, before=None):
        if not wanted(name):
            dialog.deleteLater()
            return
        dialog.setAttribute(Qt.WidgetAttribute.WA_DontShowOnScreen, True)
        dialog.show()
        deadline = time.monotonic() + 15
        while ready is not None and not ready() and time.monotonic() < deadline:
            pump(0.1)
        if before is not None:
            before(dialog)
        pump(0.4)
        dialog.grab().save(str(out / f"{name}.png"))
        print("saved", name)
        dialog.close()
        dialog.deleteLater()

    def focus_shot(name: str, top: QWidget, widget: QWidget):
        """鍵盤焦點：整張＋焦點元件附近放大的局部。"""
        if not wanted(name):
            return
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", DeprecationWarning)
            QApplication.setActiveWindow(top)
        widget.setFocus(Qt.FocusReason.TabFocusReason)
        pump(0.5)
        top.grab().save(str(out / f"{name}.png"))
        origin = widget.mapTo(top, QPoint(0, 0))
        area = QRect(origin.x() - 60, origin.y() - 40, widget.width() + 120, widget.height() + 80)
        crop = top.grab(area.intersected(top.rect()))
        crop.scaled(crop.size() * 2, Qt.AspectRatioMode.KeepAspectRatio,
                    Qt.TransformationMode.SmoothTransformation).save(str(out / f"{name}_crop.png"))
        print("saved", name, "focus:", widget.hasFocus())

    # ======================================================================== 第一次使用
    fresh = work / "fresh"
    DataStore(fresh).ensure_ready(create=True)
    fresh_local = LocalConfigStore(work / "fresh-local")
    fresh_ctx = AppContext(fresh_local, fresh, AuditLog(work / "fresh-local" / "audit.log"),
                           hpdcs_factory=lambda _c: FakeHpdcs(configured=False))
    if wanted("q00"):
        fresh_window = open_window(fresh_ctx)
        fresh_window.go("query")
        shot("q00_query_first_run_checklist", fresh_window)
        fresh_window.close()
        fresh_window.deleteLater()

    # ======================================================================== 示範資料
    store = make_demo_store(work)
    data = store.root

    fake = FakeHpdcs()
    local = LocalConfigStore(work / "local")
    local.update(data_dir=str(data), last_assessor="林美華 護理師", recent_assessors=["林美華 護理師", "陳怡君 個管師"])
    ctx = AppContext(local, data, AuditLog(work / "local" / "audit.log"), hpdcs_factory=lambda _c: fake)
    window = open_window(ctx)

    query = window.pages["query"]
    referral = window.pages["referral"]
    settings = window.pages["settings"]
    plans = ctx.active_plans()

    if wanted("q11"):
        import icope_tool.ui.query_page as query_module
        saved = query_module.nhi_card, query_module.CARD_IMPORT_ERROR
        query_module.nhi_card = None
        query_module.CARD_IMPORT_ERROR = "DLL load failed while importing _scard: 找不到指定的模組。"
        no_card = open_window(ctx)
        no_card.go("query")
        shot("q11_query_card_reader_unavailable", no_card)
        no_card.close()
        no_card.deleteLater()
        query_module.nhi_card, query_module.CARD_IMPORT_ERROR = saved

    # ======================================================================== 篩檢查詢
    window.go("query")
    shot("q01_query_idle", window)
    query.id_input.setText("A123456788")
    shot("q02_query_checksum_error", window)
    query.id_input.setText(DEMO_IDS["王小明"])
    query._request = QueryRequest(99, DEMO_IDS["王小明"], "王小明", plans, False, threading.Event())
    query._on_id_changed(query.id_input.text())
    query.loading_title.setText("正在查詢國健署系統…（已等 6 秒）")
    query.loading_text.setText("第一次查詢需要先登入，通常 5–15 秒。國健署網站很慢時，可以取消或改用網站查詢。")
    query.result_stack.setCurrentIndex(1)
    shot("q03_query_loading", window)
    query._request = None
    query._on_id_changed(query.id_input.text())

    now = datetime.now()
    can = IcopeResult([PlanResult("EFA_115", "can_assess", "今年可以繼續評估：O，可以繼續評估！"),
                       PlanResult("EFA_Pilot_115", "can_assess", "今年可以繼續評估：O，可以繼續評估！")])
    done = IcopeResult([PlanResult("EFA_115", "done", "今年可以繼續評估：X，該身分證已被登錄！[ICOPE評估表]"),
                        PlanResult("EFA_Pilot_115", "done_other", "今年無法繼續評估：X，該身分證已被其它計畫登錄！")])
    blocked = IcopeResult([PlanResult("EFA_115", "blocked", "今年無法繼續評估：X，年齡不符合本計畫收案條件！"),
                           PlanResult("EFA_Pilot_115", "blocked", "今年無法繼續評估：X，年齡不符合本計畫收案條件！")])
    ctx.add_history(HistoryEntry(now - timedelta(days=1), DEMO_IDS["張美玉"], "張美玉", can, plans=plans))
    ctx.add_history(HistoryEntry(now, DEMO_IDS["陳秀英"], "陳秀英", blocked, plans=plans))
    ctx.add_history(HistoryEntry(now, DEMO_IDS["林阿土"], "林阿土", done, plans=plans))
    ctx.set_patient(DEMO_IDS["王小明"], "王小明")
    entry = HistoryEntry(now, DEMO_IDS["王小明"], "王小明", can, plans=plans)
    ctx.add_history(entry)
    query.show_entry(entry)
    shot("q04_query_result_can_assess", window)
    query.id_input.setText(DEMO_IDS["林阿土"])
    query.show_entry(ctx.history[1])
    shot("q05_query_result_done", window)
    query.id_input.setText(DEMO_IDS["陳秀英"])
    query.show_entry(ctx.history[2])
    shot("q06_query_result_blocked", window)
    query.id_input.setText("B123456780")
    query.show_entry(entry)
    shot("q08_query_input_mismatch", window)
    query.id_input.setText(DEMO_IDS["王小明"])
    query.show_entry(entry)
    fake.disabled = True
    failing = QueryRequest(100, DEMO_IDS["王小明"], "王小明", plans, False, threading.Event())
    query._request_seq = 100
    query._on_query_error(failing, CredentialError(
        "國健署系統登入失敗（帳號密碼錯誤或帳號狀態異常），已暫停自動登入以免帳號被鎖"))
    shot("q07_query_error_login_disabled", window)
    fake.disabled = False
    query._request = QueryRequest(101, DEMO_IDS["王小明"], "王小明", plans, False, threading.Event())
    query._request_seq = 101
    query.cancel_query()
    shot("q09_query_cancelled", window)
    ctx.add_history(entry)                      # 把上面失敗那筆換回成功結果
    query.show_entry(entry)
    if wanted("q12") or wanted("q13"):
        from icope_tool.services.card.nhi_card import CardReadError
        query.card_status_text.setText("請把健保卡插入讀卡機（等待 15 秒；按 Esc 或再按一次可取消）")
        query.card_status.show()
        query.card_button.setText("取消讀卡")
        shot("q12_query_card_reading", window)
        query._card_finished()
        query._on_card_error(query._card_seq, CardReadError("NO_CARD", "讀卡機裡沒有健保卡"))
        shot("q13_query_card_error", window)
        query.card_error.hide()
    if wanted("q14"):
        query.show_entry(entry)
        query.id_input.setText(DEMO_IDS["張美玉"])          # 已經在輸入下一位
        store.update_settings(lambda s: setattr(s.hpdcs, "pilot", False))
        ctx.mark_saved("settings")                           # 查詢計畫變更：畫面上的結果過期
        shot("q14_query_result_expired_keeps_typed_id", window)
        store.update_settings(lambda s: setattr(s.hpdcs, "pilot", True))
        ctx.mark_saved("settings")
        query.id_input.setText(DEMO_IDS["王小明"])
        query.show_entry(entry)
    focus_shot("f01_focus_id_input", window, query.id_input)
    history_rows = [query.history_list.itemAt(i).widget() for i in range(query.history_list.count())]
    history_rows = [row for row in history_rows if row is not None]
    if history_rows:
        focus_shot("f02_focus_history_row", window, history_rows[0])
        shot("q10_query_history_with_stale_entry", window)

    from icope_tool.ui.dialogs.captcha import CaptchaDialog
    dialog_shot("d01_captcha", CaptchaDialog(window, CaptchaManualRequired(captcha_gif(), "image/gif", "tok"),
                                             fake.refresh_captcha))
    from icope_tool.ui.dialogs.credentials import CredentialsDialog
    dialog_shot("d02_credentials", CredentialsDialog(window, ctx))

    # ======================================================================== 轉介衛教單
    window.go("referral")
    referral.go_to(1)
    shot("r08_referral_step2_no_domains", window)
    referral.go_to(0)
    shot("r00_referral_step1_empty", window)
    referral.domain_step.cards["cognitive"].setChecked(True)
    referral.domain_step.cards["mobility"].setChecked(True)
    referral.domain_step.cards["hearing"].setChecked(True)
    shot("r01_referral_step1_selected", window)
    focus_shot("f03_focus_domain_card_checked", window, referral.domain_step.cards["cognitive"])
    referral.summary.defaults_button.click()                 # 帶入院所預設：認知、行動有預設，聽力沒有
    referral.toast.hide()
    shot("r10_referral_defaults_applied", window)
    referral.go_to(1)
    resources = {r.name: r for r in store.list_resources()}
    memory = resources["嘉義基督教醫院 記憶門診"]
    referral.state.set_resource(memory, "cognitive", True)
    station = next(r for r in store.list_resources() if r.type == "巷弄長照站(社照C據點)")
    referral.state.set_resource(station, "cognitive", True)
    care = resources["嘉義市政府衛生局附設長期照顧管理中心（嘉義市長照管理中心東區辦公室）"]
    referral.state.set_resource(care, "mobility", True)
    referral.resource_step.show_domain("cognitive")
    referral._on_changed()
    shot("r02_referral_step2_resources", window)
    selected_rows = [row for row in referral.resource_step.rows if row.checkbox.isChecked()]
    if selected_rows:
        focus_shot("f04_focus_resource_row_selected", window, selected_rows[0].checkbox)
    if referral.resource_step.tabs:
        focus_shot("f05_focus_domain_tab_current", window, referral.resource_step.tabs[0])
    referral.resource_step.show_domain("hearing")
    shot("r03_referral_step2_hearing", window)
    from icope_tool.ui.dialogs.resource_picker import ResourcePickerDialog
    dialog_shot("d03_resource_picker", ResourcePickerDialog(window, store.list_resources(), "hearing", []))
    from icope_tool.ui.dialogs.leaflet_print import LeafletPrintDialog

    def tick_two_leaflets(dialog):
        for row in dialog.rows[:2]:
            row.checkbox.setChecked(True)

    dialog_shot("d12_leaflet_print", LeafletPrintDialog(window, ctx), before=tick_two_leaflets)
    referral.go_to(2)
    referral.state.set_material(store.list_materials()[1].id, True)
    referral.material_step.rebuild()
    referral._on_changed()
    shot("r04_referral_step3_materials", window)
    focus_shot("f06_focus_step_button_current", window, referral.stepper.buttons[2])
    referral.go_to(3)
    confirm = referral.confirm_step
    deadline = time.monotonic() + 30
    while confirm._generating and time.monotonic() < deadline:
        pump(0.1)
    pump(0.6)
    shot("r05_referral_step4_preview", window)
    confirm._set_zoom(False)
    shot("r06_referral_step4_fit_width", window)
    confirm._set_zoom(True)
    if confirm._pdf_path and (wanted("p0")):
        from PySide6.QtPdf import QPdfDocument
        document = QPdfDocument()
        document.load(str(confirm._pdf_path))
        for page_index in range(min(3, document.pageCount())):
            name = f"p0{page_index + 1}_printed_page{page_index + 1}"
            document.render(page_index, QSize(1240, 1754)).save(str(out / f"{name}.png"))
            print("saved", name)
        document.close()
    if wanted("p1"):
        # 衛教重點被切在兩頁的情況（整組比一頁高時的續頁）：兩頁都要看，才看得出有沒有裁切、重複或漏字
        from PySide6.QtPdf import QPdfDocument
        from icope_tool.services.pdf.referral_sheet import ReferralJob, ReferralSection, render_referral_sheet
        fillers = [Resource(name=f"社區據點 {i + 1}", phones=["05-1234567"], address="示範市東區示範路 100 號")
                   for i in range(6)]
        long_block = Resource(name="備註很長的資源", phones=["05-7654321"], note="請先電話預約，並攜帶健保卡。" * 200)
        note = "每天做腿部肌力與平衡運動，扶著椅背練習坐下、站起；走道保持明亮、移開雜物，浴室加裝扶手。" * 4
        job = ReferralJob(clinic=store.load_settings().clinic,
                          sections=[ReferralSection("cognitive", fillers), ReferralSection("mobility", [long_block])],
                          patient_name="王小明", domain_notes={"mobility": note[:200]})
        split_pdf = work / "note-across-pages.pdf"
        split_pdf.write_bytes(render_referral_sheet(job))
        document = QPdfDocument()
        document.load(str(split_pdf))
        for page_index in range(min(2, document.pageCount())):
            name = f"p1{page_index}_note_across_pages_page{page_index + 1}"
            document.render(page_index, QSize(1240, 1754)).save(str(out / f"{name}.png"))
            print("saved", name)
        document.close()
    confirm.notes.setPlainText("回診時請帶所有藥袋與血壓紀錄本。" * 26)
    shot("r07_referral_notes_over_limit", window)
    confirm.notes.setPlainText("回診時請帶所有藥袋與血壓紀錄本；起身時放慢速度，浴室加裝扶手。")
    ctx.set_patient(DEMO_IDS["林阿土"], "林阿土", "card")
    shot("r09_referral_patient_switch_banner", window)
    referral.patient_banner.hide()

    # ======================================================================== 設定
    window.go("settings")
    for index, key in enumerate(("clinic", "hpdcs", "resources", "materials", "pack", "local"), start=1):
        settings.select(key)
        shot(f"s0{index}_settings_{key}", window)

    settings.select("hpdcs")
    focus_shot("f07_focus_checkbox", window, settings.pages["hpdcs"].official)
    focus_shot("f08_focus_settings_tabbar", window, settings.tabs.tabBar())
    settings.select("resources")
    resources_tab = settings.pages["resources"]
    resources_tab.table.selectRow(0)
    focus_shot("f09_focus_table", window, resources_tab.table)

    clinic = settings.pages["clinic"]
    settings.select("clinic")
    clinic.phone.setText("05-2761235")
    other_machine = DataStore(data)
    other_machine.update_settings(lambda s: setattr(s.clinic, "address", "嘉義市東區忠孝路 102 號"))
    later = time.time() + 5
    os.utime(data / "settings.json", (later, later))
    ctx.check_external_changes()
    shot("s07_settings_clinic_conflict", window)
    clinic.load()
    clinic.scroll_view.ensureWidgetVisible(clinic.note_count)
    shot("s09_settings_sheet_notes", window)
    clinic.note_text.setPlainText("每天做腿部肌力與平衡運動。" * 18)
    clinic.scroll_view.ensureWidgetVisible(clinic.note_count)
    shot("s10_settings_sheet_notes_over_limit", window)
    clinic.load()

    materials_tab = settings.pages["materials"]
    victim = store.list_materials()[2]
    store.material_file(victim).unlink()
    settings.select("materials")
    materials_tab.reload()
    materials_tab._select(victim.id)
    shot("s08_settings_materials_missing_file", window)

    from icope_tool.ui.dialogs.editors import MaterialDialog, ResourceDialog
    dialog_shot("d04_resource_edit", ResourceDialog(window, store.resource_types(), memory))
    dialog_shot("d05_material_edit", MaterialDialog(window, store.list_materials()[0]))
    xlsx = work / "盤點表.xlsx"
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "嘉義市1003"
    ws.append(["序號", "分類名稱", "個別名稱", "型態", "類別", "地址", "電話", "電郵", "官網"])
    ws.append([1, "巷弄長照站(社照C據點)", "嘉義市東區後湖社區發展協會", "1-1", "1", "嘉義市東區後湖里保義路60巷40號", "0935-214218"])
    ws.append([2, "聽力所", "新的聽力中心", "3-1", "3E", "嘉義市西區中山路 1 號", "05-1234567"])
    wb.save(xlsx)
    from icope_tool.ui.dialogs.importing import ExcelImportDialog, PackImportDialog
    excel = ExcelImportDialog(window, store, xlsx)
    dialog_shot("d06_excel_import", excel, ready=excel.import_button.isEnabled)
    pack = read_pack(ROOT / "examples" / "chiayi-city-resources.json")

    def choose_replace(dialog):
        dialog.replace.setChecked(True)

    dialog_shot("d07_pack_import_replace", PackImportDialog(window, store, pack), before=choose_replace)
    if wanted("f10"):
        radio_dialog = PackImportDialog(window, store, pack)
        radio_dialog.setAttribute(Qt.WidgetAttribute.WA_DontShowOnScreen, True)
        radio_dialog.show()
        pump(0.3)
        focus_shot("f10_focus_radio", radio_dialog, radio_dialog.findChildren(QRadioButton)[0])
        radio_dialog.close()
        radio_dialog.deleteLater()
    from icope_tool.ui.dialogs.data_folder import DataFolderDialog
    dialog_shot("d08_first_run", DataFolderDialog(None, "first_run"))
    dialog_shot("d09_data_unavailable", DataFolderDialog(None, "unavailable", Path(r"\\NAS\icope\data")))
    from icope_tool.ui.help_dialog import HelpDialog
    dialog_shot("d10_help", HelpDialog(window))

    def scroll_to_end(dialog):
        from PySide6.QtWidgets import QTextBrowser
        browser = dialog.findChild(QTextBrowser)
        browser.verticalScrollBar().setValue(browser.verticalScrollBar().maximum())

    dialog_shot("d11_help_shortcuts", HelpDialog(window), before=scroll_to_end)

    referral.confirm_step.release_document()        # 關閉時不要跳「轉介單還沒印出」
    window.close()
    shutil.rmtree(work, ignore_errors=True)
    return 0


if __name__ == "__main__":
    os.environ.setdefault("QT_SCALE_FACTOR", "1")
    raise SystemExit(main())
