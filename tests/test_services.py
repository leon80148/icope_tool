"""匯入、設定包、PDF 產生的測試。"""
from __future__ import annotations

import zipfile
from datetime import date
from pathlib import Path

import openpyxl
import pytest
from pypdf import PdfReader

from icope_tool.models import ClinicInfo, Material, Resource
from icope_tool.services import importer, pack
from icope_tool.services.pdf.referral_sheet import (
    LeafletSource, MaterialChanged, PdfBuildError, ReferralJob, ReferralSection, build_referral_pdf,
    merge_material_pdfs, render_referral_sheet, verify_material_sources,
)
from icope_tool.store import DataStore, sha256_file
from tests.test_store import make_pdf

EXAMPLES = Path(__file__).resolve().parent.parent / "examples"


@pytest.fixture
def store(tmp_path) -> DataStore:
    s = DataStore(tmp_path / "data")
    s.ensure_ready(create=True)
    return s


# ---------------------------------------------------------------------------
# Excel 匯入
# ---------------------------------------------------------------------------
def _inventory_xlsx(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "舊資料"
    ws2 = wb.create_sheet("嘉義市1003")
    ws2.append(["嘉義市社區資源盤點表"])
    ws2.append(["序號", "分類名稱", "個別名稱", "型態", "類別", "地址", "電話", "電郵", "官網", "其他"])
    ws2.append(["範例", "外縣市", "某某縣資源（範例列，應略過）", "1-1", "1A", "某縣", "000", None, None, None])
    ws2.append([1, "巷弄長照站(社照C據點)", "嘉義市東區後湖社區發展協會", "1-1", "1AB", "嘉義市東區後湖里保義路60巷40號", "0935-214218", None, "-", None])
    ws2.append([2, None, "延續列沿用上一列分類", "2-1", "2", "嘉義市西區", "05-1111111、0912-345678", None, "https://x.example", None])
    ws2.append([3, "聽力所", "好市多聽力中心-嘉義店", "3-1", "3E", "嘉義市東區忠孝路668號", "05-3200787", None, None, None])
    wb.save(path)
    return path


def test_parse_inventory_format_picks_sheet_and_maps_domains(tmp_path):
    preview = importer.parse_resources_xlsx(_inventory_xlsx(tmp_path / "inv.xlsx"))
    assert preview.format == "inventory" and preview.sheet == "嘉義市1003"
    assert [r.name for r in preview.resources] == [
        "嘉義市東區後湖社區發展協會", "延續列沿用上一列分類", "好市多聽力中心-嘉義店"]
    first, second, third = preview.resources
    assert first.domains == ["cognitive", "mobility"] and first.type == "巷弄長照站(社照C據點)"
    assert first.website == ""
    assert second.type == "巷弄長照站(社照C據點)" and second.phones == ["05-1111111", "0912-345678"]
    assert third.domains == ["hearing"]
    assert any("沒有 A–F" in w for w in preview.warnings)


def test_template_roundtrip(tmp_path):
    path = tmp_path / "template.xlsx"
    importer.write_template_xlsx(path)
    wb = openpyxl.load_workbook(path)
    ws = wb["轉介資源"]
    ws.append(["嘉義市輔具資源中心", "輔具資源中心", "行動、聽力", "05-2256686", "嘉義市東區彌陀路255號",
               "王小姐", "週一至週五", "", "", "是"])
    ws.append(["全部都適用", "門診追蹤", "全部", "", "", "", "", "", "", ""])
    ws.append(["看不懂的項目", "", "腸胃", "", "", "", "", "", "", ""])
    ws.append([None, "只有類型沒有名稱", "", "", "", "", "", "", "", ""])
    wb.save(path)

    preview = importer.parse_resources_xlsx(path)
    assert preview.format == "template"
    names = [r.name for r in preview.resources]
    assert names == ["嘉義市輔具資源中心", "全部都適用", "看不懂的項目"]   # 範例列略過
    assistive = preview.resources[0]
    assert assistive.domains == ["mobility", "hearing"] and assistive.pinned and assistive.contact == "王小姐"
    assert len(preview.resources[1].domains) == 8
    assert preview.skipped_rows == 1
    assert any("腸胃" in w for w in preview.warnings)


def test_unrecognized_sheet_raises(tmp_path):
    path = tmp_path / "x.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["隨便", "的", "表頭"])
    wb.save(path)
    with pytest.raises(importer.ImportError_):
        importer.parse_resources_xlsx(path)


def test_not_excel_raises(tmp_path):
    path = tmp_path / "x.xlsx"
    path.write_text("not excel")
    with pytest.raises(importer.ImportError_):
        importer.parse_resources_xlsx(path)


# ---------------------------------------------------------------------------
# 設定包
# ---------------------------------------------------------------------------
def test_pack_roundtrip_merge_and_replace(store, tmp_path):
    store.add_resource(Resource(name="輔具中心", address="彌陀路255號", domains=["mobility"]))
    store.add_material_file(make_pdf(tmp_path / "營養.pdf", pages=2), domains=["nutrition"])
    (store.auth_dir / "hpdcs_credentials.json").write_text('{"account":"a","password":"secret"}')

    dest = tmp_path / "pack.zip"
    assert pack.export_pack(store, dest) == (2, 1)
    with zipfile.ZipFile(dest) as archive:
        names = archive.namelist()
        assert "pack.json" in names and all(not n.startswith("auth") for n in names)
        assert b"secret" not in archive.read("pack.json")

    other = DataStore(tmp_path / "other")
    other.ensure_ready(create=True)
    loaded = pack.read_pack(dest)
    pack.validate_pack_materials(loaded, tmp_path / "scratch")
    result = pack.apply_pack(other, loaded, "merge")
    assert (result.resources_added, result.resources_skipped) == (1, 1)   # 預設「本院門診追蹤」重複
    assert result.materials_added == 1
    again = pack.apply_pack(other, pack.read_pack(dest), "merge")
    assert again.resources_added == 0 and again.materials_skipped == 1

    third = DataStore(tmp_path / "third")
    third.ensure_ready(create=True)
    third.add_resource(Resource(name="會被取代掉"))
    replaced = pack.apply_pack(third, pack.read_pack(dest), "replace")
    assert replaced.resources_added == 2 and replaced.materials_added == 1
    assert {r.name for r in third.list_resources()} == {"本院門診追蹤", "輔具中心"}
    material = third.list_materials()[0]
    assert third.material_file(material).exists() and material.pages == 2


def test_pack_carries_default_flags(store, tmp_path):
    store.add_resource(Resource(name="輔具中心", address="彌陀路255號", domains=["mobility"], include_by_default=True))
    store.add_material_file(make_pdf(tmp_path / "營養.pdf"), domains=["nutrition"], include_by_default=True)
    dest = tmp_path / "pack.zip"
    pack.export_pack(store, dest)

    other = DataStore(tmp_path / "other")
    other.ensure_ready(create=True)
    other.add_resource(Resource(name="輔具中心", address="彌陀路255號"))          # 本院已有這筆，沒有設為預設
    pack.apply_pack(other, pack.read_pack(dest), "merge")
    local = [r for r in other.list_resources() if r.name == "輔具中心"]
    assert len(local) == 1 and local[0].include_by_default is False                # 重複：保留本院的設定
    assert other.list_materials()[0].include_by_default is True                    # 新增的單張帶著旗標

    third = DataStore(tmp_path / "third")
    third.ensure_ready(create=True)
    pack.apply_pack(third, pack.read_pack(dest), "replace")
    assert next(r for r in third.list_resources() if r.name == "輔具中心").include_by_default is True
    assert third.list_materials()[0].include_by_default is True


@pytest.mark.parametrize("mode", ["merge", "replace"])
def test_pack_import_never_touches_the_clinic_written_sheet_text(store, tmp_path, mode):
    """院所衛教文字存在共用資料夾、不在設定包裡：兩種匯入模式都不能動它。"""
    source = DataStore(tmp_path / "source")
    source.ensure_ready(create=True)
    source.update_settings(lambda s: s.sheet.domain_notes.update({"cognitive": "別家診所寫的"}))
    source.add_resource(Resource(name="輔具中心", domains=["mobility"]))
    dest = tmp_path / "pack.zip"
    pack.export_pack(source, dest)
    with zipfile.ZipFile(dest) as archive:
        assert "別家診所寫的" not in archive.read("pack.json").decode("utf-8")

    store.update_settings(lambda s: s.sheet.domain_notes.update({"cognitive": "本院寫的衛教重點"}))
    pack.apply_pack(store, pack.read_pack(dest), mode)
    assert store.load_settings().sheet.domain_notes == {"cognitive": "本院寫的衛教重點"}


def test_old_packs_and_excel_never_set_the_default_flag(tmp_path):
    loaded = pack.read_pack(EXAMPLES / "chiayi-city-resources.json")
    assert loaded.resource_count > 200 and not any(r.include_by_default for r in loaded.manifest.resources)
    preview = importer.parse_resources_xlsx(_inventory_xlsx(tmp_path / "inv.xlsx"))
    assert preview.resources and not any(r.include_by_default for r in preview.resources)


def test_read_pack_rejects_foreign_files(tmp_path):
    bad_zip = tmp_path / "bad.zip"
    with zipfile.ZipFile(bad_zip, "w") as archive:
        archive.writestr("readme.txt", "hi")
    with pytest.raises(pack.PackError):
        pack.read_pack(bad_zip)
    not_zip = tmp_path / "x.zip"
    not_zip.write_text("nope")
    with pytest.raises(pack.PackError):
        pack.read_pack(not_zip)
    wrong_format = tmp_path / "x.json"
    wrong_format.write_text('{"format": "other"}', encoding="utf-8")
    with pytest.raises(pack.PackError):
        pack.read_pack(wrong_format)


def test_read_pack_refuses_to_inflate_an_oversized_pack(store, tmp_path, monkeypatch):
    """設定包會在診所之間流傳：宣告的解壓後大小超過上限就不讀進記憶體（壓縮炸彈只要幾 KB 就能佔滿記憶體）。"""
    bomb = tmp_path / "bomb.zip"
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("pack.json", '{"format": "icope-tool-pack", "pad": "' + "0" * (pack.MAX_MANIFEST_BYTES + 1) + '"}')
    assert bomb.stat().st_size < 64 * 1024
    with pytest.raises(pack.PackError, match="過大"):
        pack.read_pack(bomb)

    for index in range(3):
        store.add_material_file(make_pdf(tmp_path / f"{index}.pdf", text=f"leaflet {index}"), name=f"單張 {index}")
    exported = tmp_path / "three.zip"
    pack.export_pack(store, exported)
    assert pack.read_pack(exported).material_count == 3
    one = max(m.size for m in store.list_materials())
    monkeypatch.setattr(pack, "MAX_PACK_BYTES", one * 2)             # 三份加起來超過上限
    with pytest.raises(pack.PackError, match="過大"):
        pack.read_pack(exported)


def test_example_resources_json_is_importable(store):
    loaded = pack.read_pack(EXAMPLES / "chiayi-city-resources.json")
    assert loaded.resource_count > 200 and loaded.material_count == 0
    result = pack.apply_pack(store, loaded, "merge")
    assert result.resources_added > 200


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def _job(**overrides) -> ReferralJob:
    base = dict(
        clinic=ClinicInfo(name="測試診所", phone="05-0000000", address="嘉義市", footer_note="僅供參考"),
        sections=[
            ReferralSection("cognitive", [Resource(name="記憶門診", type="醫療院所", phones=["05-1"],
                                                   address="很長的地址" * 12, note="需先掛號")]),
            ReferralSection("hearing", []),
        ],
        patient_name="王小明", assessor="林護理師", notes="回診請帶藥袋", assessed_on=date(2026, 9, 17),
    )
    base.update(overrides)
    return ReferralJob(**base)


def test_render_sheet_contains_chinese_text():
    reader = PdfReader(__import__("io").BytesIO(render_referral_sheet(_job())))
    text = "".join(page.extract_text() for page in reader.pages)
    for expected in ("測試診所", "王小明", "記憶門診", "民國 115 年 9 月 17 日", "聽力"):
        assert expected in text


def test_long_content_flows_to_more_pages():
    many = [Resource(name=f"資源 {i}", phones=["05-1234567"], address="嘉義市東區忠孝路539號", note="備註" * 10)
            for i in range(40)]
    reader = PdfReader(__import__("io").BytesIO(render_referral_sheet(_job(sections=[ReferralSection("mobility", many)]))))
    assert len(reader.pages) >= 3


def test_build_merges_materials(tmp_path):
    m = Material(id="aaaaaaaaaaaa", name="營養", filename="aaaaaaaaaaaa.pdf", pages=3)
    files = {m.id: make_pdf(tmp_path / "m.pdf", pages=3)}
    result = build_referral_pdf(_job(materials=[m]), files, tmp_path / "out" / "sheet.pdf")
    assert result.total_pages == result.sheet_pages + 3
    assert len(PdfReader(str(result.path)).pages) == result.total_pages


def test_build_reports_missing_material(tmp_path):
    m = Material(id="bbbbbbbbbbbb", name="已刪除的單張", filename="x.pdf")
    with pytest.raises(PdfBuildError) as ei:
        build_referral_pdf(_job(materials=[m]), {}, tmp_path / "o.pdf")
    assert "已刪除的單張" in str(ei.value)


def test_materials_only_sheet(tmp_path):
    m = Material(id="cccccccccccc", name="運動", filename="c.pdf", pages=1)
    result = build_referral_pdf(_job(sections=[], patient_name="", assessor="", notes="", materials=[m]),
                                {m.id: make_pdf(tmp_path / "c.pdf")}, tmp_path / "o.pdf")
    assert result.sheet_pages == 1


# ---------------------------------------------------------------------------
# 各項衛教重點（院所自己寫、印在轉介單各項目標題下）
# ---------------------------------------------------------------------------
def _pdf_fragments(data: bytes) -> list[list[tuple[str, float, float]]]:
    """每頁的（文字, x, y）；座標是 PDF 的點，原點在左下角。"""
    pages = []
    for page in PdfReader(__import__("io").BytesIO(data)).pages:
        fragments: list[tuple[str, float, float]] = []

        def visit(text, _cm, tm, _font, _size, out=fragments):
            if text.strip():
                out.append((text, float(tm[4]), float(tm[5])))

        page.extract_text(visitor_text=visit)
        pages.append(fragments)
    return pages


def test_sheet_prefs_tolerate_hand_edited_values_without_losing_text():
    from icope_tool.models import NOTE_MAX_CHARS, Settings, SheetPrefs, note_problem
    assert Settings.model_validate({}).sheet.domain_notes == {}                  # 舊的 settings.json
    assert Settings.model_validate({"sheet": "壞掉了"}).sheet.domain_notes == {}
    loaded = Settings.model_validate({"sheet": {"domain_notes": {
        "hearing": "  聽不清楚請告訴家人\r\n可以到聽力所檢查  ", "unknown": "不存在的項目", "vision": 5, "social": "   "}}})
    assert loaded.sheet.domain_notes == {"hearing": "聽不清楚請告訴家人\n可以到聽力所檢查"}
    kept = Settings(sheet=SheetPrefs(domain_notes={"cognitive": "多動腦"}))       # 程式內正常建立的物件不能被容錯清掉
    assert kept.sheet.domain_notes == {"cognitive": "多動腦"}
    too_long = "字" * (NOTE_MAX_CHARS + 1)
    assert Settings.model_validate({"sheet": {"domain_notes": {"mobility": too_long}}}).sheet.domain_notes == {
        "mobility": too_long}                                                    # 超限：保留不截斷，由 note_problem 回報
    assert note_problem("字" * NOTE_MAX_CHARS) == "" and "200" in note_problem(too_long)
    assert note_problem("一\n二\n三\n四\n五\n六") == "" and "6" in note_problem("一\n二\n三\n四\n五\n六\n七")
    assert note_problem("\r\n".join("字" * 40 for _ in range(5)) + "  \r\n") == ""   # 換行與首尾空白不算字數


def test_domain_notes_print_only_for_selected_domains(tmp_path):
    notes = {"cognitive": "多動腦、多互動，每天和家人聊聊今天做了什麼。", "hearing": "聽不清楚時請告訴家人。",
             "vision": "沒有勾選的項目不會印"}
    text = "".join(p.extract_text() for p in PdfReader(__import__("io").BytesIO(
        render_referral_sheet(_job(domain_notes=notes)))).pages)
    assert "多動腦、多互動" in text and "沒有勾選的項目不會印" not in text
    assert "聽不清楚時請告訴家人" in text and "請與本院醫護人員討論後續安排" in text      # 沒有資源也照印，原本的句子保留
    assert text.index("多動腦") < text.index("記憶門診")                              # 在資源之前
    assert text.count("多動腦、多互動") == 1
    plain = "".join(p.extract_text() for p in PdfReader(__import__("io").BytesIO(render_referral_sheet(_job()))).pages)
    assert "多動腦" not in plain                                                      # 沒填就不占位置


def test_over_limit_note_blocks_only_sheets_that_include_that_domain():
    too_long = {"vision": "字" * 201}
    render_referral_sheet(_job(domain_notes=too_long))                                # 這份沒有視力：不受影響
    with pytest.raises(PdfBuildError) as ei:
        render_referral_sheet(_job(sections=[ReferralSection("vision", [])], domain_notes=too_long))
    assert "視力" in str(ei.value) and "診所與轉介單" in str(ei.value)


@pytest.mark.parametrize("filler", [9, 10, 11, 12])
def test_domain_notes_paginate_without_clipping(filler):
    """長文字落在頁尾附近、多個項目同時有文字：內容完整、不重複、不畫進頁尾、沒有空白頁。"""
    note = "".join(f"第{i}句衛教重點要寫得讓長者看得懂。" for i in range(1, 12))[:200]
    resources = [Resource(name=f"填充資源 {i}", phones=["05-1234567"], address="示範市東區示範路 100 號")
                 for i in range(filler)]
    job = _job(sections=[ReferralSection("cognitive", resources), ReferralSection("mobility", resources[:1]),
                         ReferralSection("hearing", [])],
               domain_notes={"mobility": note, "hearing": note.replace("衛教", "聽力")})
    data = render_referral_sheet(job)
    pages = _pdf_fragments(data)
    whole = "".join(text for page in pages for text, _x, _y in page)
    assert whole.count("第1句衛教重點") == 1 and whole.count("第1句聽力重點") == 1
    for i in range(1, 9):
        assert f"第{i}句衛教重點" in whole                                           # 沒有漏字
    assert len(pages) >= 2 and all(page for page in pages)                           # 沒有空白頁
    footer_top = 22 / 25.4 * 72                                                      # 下邊界 22 mm
    for page in pages:
        body = [(t, x, y) for t, x, y in page if "重點" in t]
        assert all(y > footer_top and x >= 16 / 25.4 * 72 - 1 for _t, x, y in body)   # 不畫進頁尾、不超出左邊界


def test_note_plus_first_resource_taller_than_a_page_still_flows():
    """資源備註沒有長度上限：「標題＋衛教重點＋第一筆資源」比一頁還高時改成正常續頁，衛教重點可能被切在兩頁。"""
    huge = Resource(name="備註非常長的資源", phones=["05-1"], note="請先電話預約。" * 400)
    note = "".join(f"第{i}句衛教重點要寫得讓長者看得懂。" for i in range(1, 12))[:200]
    top, footer_top = (297 - 16 - 10) / 25.4 * 72, 22 / 25.4 * 72                    # 續頁頁首以下、頁尾以上
    straddled = False
    for filler in range(4, 10):                                                      # 讓這一段落在頁面上不同的高度
        resources = [Resource(name=f"填充資源 {i}", phones=["05-1234567"], address="示範市東區示範路 100 號")
                     for i in range(filler)]
        job = _job(sections=[ReferralSection("cognitive", resources), ReferralSection("mobility", [huge])],
                   domain_notes={"mobility": note})
        pages = _pdf_fragments(render_referral_sheet(job))
        assert len(pages) >= 3 and all(page for page in pages)                       # 正常續頁、沒有空白頁
        whole = "".join(text for page in pages for text, _x, _y in page)
        assert all(whole.count(f"第{i}句衛教重點") == 1 for i in range(1, 9))         # 完整、不重複
        title_page = next(i for i, page in enumerate(pages) if any("行動能力" in t for t, _x, _y in page))
        assert any("第1句衛教重點" in t for t, _x, _y in pages[title_page])           # 標題和第一行同頁
        with_note = [[(x, y) for t, x, y in page if "衛教重點" in t] for page in pages]
        for index, rows in enumerate(with_note):
            assert all(footer_top < y < (top if index else 842) for _x, y in rows)   # 不進頁尾、續頁不壓到頁首
        assert len({round(x, 1) for rows in with_note for x, _y in rows}) == 1       # 跨頁後縮排一致
        for page in pages[:-1]:                                                      # 除了最後一頁，每頁都印到接近下邊界
            assert min(y for _t, _x, y in page if y > footer_top) < 130               # 沒有「只印幾行、其餘空白」的頁面
        straddled = straddled or sum(1 for rows in with_note if rows) == 2
    assert straddled                                                                 # 確實測到「文字被切在兩頁」


# ---------------------------------------------------------------------------
# 只印衛教單張：合併與來源驗證
# ---------------------------------------------------------------------------
def _plain_pdf(path: Path, text: str) -> Path:
    """不壓縮、固定建立時間：同樣長度的文字會得到同樣長度的檔案（用來做「同長度、不同內容」的反例）。"""
    from datetime import datetime, timezone

    from fpdf import FPDF
    pdf = FPDF()
    pdf.set_compression(False)
    pdf.creation_date = datetime(2026, 1, 1, tzinfo=timezone.utc)
    pdf.add_page()
    pdf.set_font("helvetica", size=12)
    pdf.cell(text=text)
    pdf.output(str(path))
    return path


def test_leaflet_merge_has_no_cover_and_preserves_order(tmp_path):
    first = make_pdf(tmp_path / "b.pdf", pages=1, text="beta")
    second = make_pdf(tmp_path / "a.pdf", pages=2, text="alpha")
    sources = [LeafletSource("b", "乙", first, sha256_file(first)), LeafletSource("a", "甲", second, sha256_file(second))]
    result = merge_material_pdfs(sources, tmp_path / "out" / "merged.pdf")
    texts = [page.extract_text() for page in PdfReader(str(result.path)).pages]
    assert result.pages == 3 == len(texts)                       # 頁數來自合併結果
    assert "beta" in texts[0] and "alpha" in texts[1] and "alpha" in texts[2]
    assert not any("轉介" in text or "測試診所" in text for text in texts)     # 沒有轉介單首頁
    assert result.hashes == (sha256_file(first), sha256_file(second))


def test_leaflet_merge_refuses_empty_missing_and_corrupt_sources(tmp_path):
    out = tmp_path / "merged.pdf"
    with pytest.raises(PdfBuildError):
        merge_material_pdfs([], out)
    good = make_pdf(tmp_path / "good.pdf")
    with pytest.raises(PdfBuildError) as missing:
        merge_material_pdfs([LeafletSource("g", "好的", good, ""), LeafletSource("m", "不見的單張", tmp_path / "no.pdf", "")], out)
    assert "不見的單張" in str(missing.value)
    broken = tmp_path / "broken.pdf"
    broken.write_bytes(b"%PDF-1.4 not really")
    with pytest.raises(PdfBuildError) as corrupt:
        merge_material_pdfs([LeafletSource("x", "壞掉的單張", broken, "")], out)
    assert "壞掉的單張" in str(corrupt.value)
    assert not out.exists()                                       # 失敗不留半成品


def test_leaflet_merge_rejects_a_source_that_no_longer_matches_its_registered_hash(tmp_path):
    source = make_pdf(tmp_path / "a.pdf", text="registered")
    registered = sha256_file(source)
    source.write_bytes(make_pdf(tmp_path / "other.pdf", text="swapped").read_bytes())
    with pytest.raises(MaterialChanged):
        merge_material_pdfs([LeafletSource("a", "甲", source, registered)], tmp_path / "merged.pdf")
    assert not (tmp_path / "merged.pdf").exists()


def test_same_length_swap_with_restored_mtime_is_still_detected(tmp_path):
    """第五輪反例：沒有登錄雜湊、檔案長度與修改時間都一樣、JSON 沒動，內容換了也要抓得到。"""
    import os
    source = _plain_pdf(tmp_path / "a.pdf", "fall prevention A")
    other = _plain_pdf(tmp_path / "other.pdf", "fall prevention B")
    assert source.stat().st_size == other.stat().st_size and source.read_bytes() != other.read_bytes()
    leaflet = LeafletSource("a", "甲", source, "")               # 沒有登錄雜湊：以實際內容建立這次的基準
    result = merge_material_pdfs([leaflet], tmp_path / "merged.pdf")
    verify_material_sources([leaflet], result.hashes)            # 沒變：通過
    stat = source.stat()
    source.write_bytes(other.read_bytes())
    os.utime(source, ns=(stat.st_atime_ns, stat.st_mtime_ns))
    assert source.stat().st_mtime_ns == stat.st_mtime_ns and source.stat().st_size == stat.st_size
    with pytest.raises(MaterialChanged):
        verify_material_sources([leaflet], result.hashes)


def test_shared_resource_printed_once_then_referenced():
    shared = Resource(name="後湖社區發展協會", phones=["0935-214218"], address="保義路60巷40號")
    job = _job(sections=[ReferralSection("cognitive", [shared]), ReferralSection("mobility", [shared])])
    reader = PdfReader(__import__("io").BytesIO(render_referral_sheet(job)))
    text = "".join(page.extract_text() for page in reader.pages)
    assert text.count("0935-214218") == 2          # 每個項目都看得到電話
    assert text.count("保義路60巷40號") == 1        # 地址等完整資訊只印一次
    assert "見上方「認知功能」" in text


def test_shared_resource_reference_names_the_page_when_split():
    shared = Resource(name="後湖社區發展協會", phones=["0935-214218"], address="保義路60巷40號")
    fillers = [Resource(name=f"填充資源 {i}", phones=["05-1234567"], address="嘉義市東區", hours="週一至週五",
                        note="需要先電話預約") for i in range(14)]
    job = _job(sections=[ReferralSection("cognitive", [shared, *fillers]), ReferralSection("mobility", [shared])])
    reader = PdfReader(__import__("io").BytesIO(render_referral_sheet(job)))
    assert len(reader.pages) >= 2
    last = reader.pages[-1].extract_text()
    assert "見第 1 頁「認知功能」" in last


def test_clinic_contact_block_printed_large_when_phone_known():
    job = _job()
    reader = PdfReader(__import__("io").BytesIO(render_referral_sheet(job)))
    text = "".join(page.extract_text() for page in reader.pages)
    assert "有問題請聯絡本院" in text and "電話 05-0000000" in text


def test_copy_data_dir_cancel_removes_partial_copy(tmp_path):
    """U24：複製資料夾中途取消，新位置不留半套資料。"""
    import threading

    from icope_tool.ui.dialogs.data_folder import CopyCancelled, copy_data_dir, data_files
    source = DataStore(tmp_path / "src")
    source.ensure_ready(create=True)
    for index in range(3):
        source.add_material_file(make_pdf(tmp_path / f"m{index}.pdf", text=f"m{index}"))
    assert len(data_files(source.root)) >= 5
    target = tmp_path / "dst"
    cancel = threading.Event()
    progress = []

    def report(index, text):
        progress.append(text)
        if index == 2:
            cancel.set()

    with pytest.raises(CopyCancelled):
        copy_data_dir(source.root, target, report, cancel)
    assert progress and not DataStore.looks_initialized(target)
    assert not any(path.is_file() for path in target.rglob("*"))


def test_copy_data_dir_cancel_during_last_file_still_cleans_up(tmp_path):
    """U24：最後一個檔案複製期間按取消，也不能留下資料或切換位置。"""
    import threading

    from icope_tool.ui.dialogs.data_folder import CopyCancelled, copy_data_dir, data_files
    source = DataStore(tmp_path / "src")
    source.ensure_ready(create=True)
    source.add_material_file(make_pdf(tmp_path / "m.pdf"))
    total = len(data_files(source.root))
    target = tmp_path / "dst"
    cancel = threading.Event()

    def report(index, _text):
        if index == total - 1:
            cancel.set()

    with pytest.raises(CopyCancelled):
        copy_data_dir(source.root, target, report, cancel)
    assert not any(path.is_file() for path in target.rglob("*"))
